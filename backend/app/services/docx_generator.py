import io
import datetime
from typing import List, Optional, Dict, Any
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from app.schemas.meeting import MeetingDebrief


# Corporate B2B Color Constants
COLOR_NAVY_HEX = "1E293B"       # Deep Slate Navy
COLOR_NAVY = RGBColor(30, 41, 59)
COLOR_INDIGO_HEX = "4F46E5"     # Precision Indigo
COLOR_INDIGO = RGBColor(79, 70, 229)
COLOR_SLATE_DARK = RGBColor(51, 65, 85)   # #334155
COLOR_SLATE_MUTED = RGBColor(100, 116, 139) # #64748B
COLOR_BORDER_HEX = "E2E8F0"     # Light Border
COLOR_BG_ALT_HEX = "F8FAFC"     # Alternating Table Row
COLOR_WHITE_HEX = "FFFFFF"


def _set_cell_background(cell, hex_color: str):
    """Sets cell background shading."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Sets internal cell margins in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _set_table_borders(table, border_color_hex: str = COLOR_BORDER_HEX):
    """Applies subtle horizontal borders with no vertical borders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color_hex}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color_hex}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color_hex}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def generate_executive_docx(
    debrief: MeetingDebrief,
    transcript_history: Optional[List[Dict[str, Any]]] = None
) -> io.BytesIO:
    """
    Constructs an executive corporate intelligence document matching corporate B2B standards.
    """
    doc = docx.Document()

    # Configure Margins (0.75 in / 54 pt)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # 1. Executive Banner Header
    banner_table = doc.add_table(rows=1, cols=1)
    banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_cell = banner_table.cell(0, 0)
    banner_cell.width = Inches(7.0)
    _set_cell_background(banner_cell, COLOR_NAVY_HEX)
    _set_cell_margins(banner_cell, top=300, bottom=300, left=300, right=300)

    bp = banner_cell.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    b_tag = bp.add_run("VOCALSYNC AI • ENTERPRISE MEETING INTELLIGENCE\n")
    b_tag.font.name = "Arial"
    b_tag.font.size = Pt(9)
    b_tag.font.bold = True
    b_tag.font.color.rgb = RGBColor(148, 163, 184)  # Slate 400

    b_title = bp.add_run(debrief.meeting_title or "Executive Meeting Intelligence Brief")
    b_title.font.name = "Arial"
    b_title.font.size = Pt(18)
    b_title.font.bold = True
    b_title.font.color.rgb = RGBColor(255, 255, 255)

    bp2 = banner_cell.add_paragraph()
    timestamp_str = datetime.datetime.now().strftime("%B %d, %Y • %I:%M %p")
    b_meta = bp2.add_run(f"Generated on {timestamp_str} | SpeechBrain ECAPA-TDNN & Faster-Whisper Biometrics")
    b_meta.font.name = "Arial"
    b_meta.font.size = Pt(8.5)
    b_meta.font.color.rgb = RGBColor(203, 213, 225)  # Slate 300

    doc.add_paragraph()  # Spacing

    # Helper for Section Headers
    def add_section_header(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY
        return p

    # 2. Executive Summary & Meeting Objective
    add_section_header("1. Meeting Objective & Scope")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_after = Pt(8)
    r_obj = p_obj.add_run(debrief.objective or "No explicit objective logged.")
    r_obj.font.name = "Calibri"
    r_obj.font.size = Pt(10.5)
    r_obj.font.color.rgb = COLOR_SLATE_DARK

    add_section_header("2. Executive Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(12)
    r_sum = p_sum.add_run(debrief.executive_summary or "Summary not available.")
    r_sum.font.name = "Calibri"
    r_sum.font.size = Pt(10.5)
    r_sum.font.color.rgb = COLOR_SLATE_DARK

    # 3. Per-Speaker Contribution & Intent Table
    add_section_header("3. Per-Speaker Intent & Contribution Matrix")
    if debrief.speaker_intents:
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(table)

        col_widths = [Inches(1.5), Inches(1.2), Inches(2.3), Inches(2.0)]
        headers = ["Identified Speaker", "Inferred Role", "Key Discussion Points", "Underlying Intent / Motive"]

        hdr_row = table.rows[0]
        for idx, (cell, h_text) in enumerate(zip(hdr_row.cells, headers)):
            cell.width = col_widths[idx]
            _set_cell_background(cell, COLOR_NAVY_HEX)
            _set_cell_margins(cell, top=160, bottom=160, left=160, right=160)
            hp = cell.paragraphs[0]
            hr = hp.add_run(h_text)
            hr.font.name = "Arial"
            hr.font.size = Pt(9.5)
            hr.font.bold = True
            hr.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, item in enumerate(debrief.speaker_intents):
            row = table.add_row()
            bg_color = COLOR_BG_ALT_HEX if (row_idx % 2 == 1) else COLOR_WHITE_HEX
            cells_data = [
                item.speaker,
                item.role_or_identity,
                item.key_contributions,
                item.underlying_intent
            ]

            for idx, (cell, val) in enumerate(zip(row.cells, cells_data)):
                cell.width = col_widths[idx]
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                p = cell.paragraphs[0]
                run = p.add_run(val)
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.color.rgb = COLOR_SLATE_DARK
                if idx == 0:
                    run.font.bold = True
    else:
        p_none = doc.add_paragraph("No speaker intent analysis logged.")
        p_none.runs[0].font.color.rgb = COLOR_SLATE_MUTED

    doc.add_paragraph()  # Spacing

    # 4. Key Decisions & Agreements Made
    add_section_header("4. Key Decisions & Strategic Alignments")
    if debrief.key_decisions:
        for dec in debrief.key_decisions:
            dp = doc.add_paragraph()
            dp.paragraph_format.left_indent = Inches(0.25)
            dp.paragraph_format.space_after = Pt(4)
            bullet = dp.add_run("✔  ")
            bullet.font.bold = True
            bullet.font.color.rgb = COLOR_INDIGO
            dr = dp.add_run(dec)
            dr.font.name = "Calibri"
            dr.font.size = Pt(10)
            dr.font.color.rgb = COLOR_SLATE_DARK
    else:
        p_none = doc.add_paragraph("No explicit decisions recorded.")
        p_none.runs[0].font.color.rgb = COLOR_SLATE_MUTED

    doc.add_paragraph()  # Spacing

    # 5. Action Items Checklist
    add_section_header("5. Action Items & Next Deliverables")
    if debrief.action_items:
        act_table = doc.add_table(rows=1, cols=3)
        act_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(act_table)

        act_widths = [Inches(4.2), Inches(1.5), Inches(1.3)]
        act_headers = ["Action Item / Task Description", "Assigned Owner", "Target Deadline"]

        for idx, (cell, h_text) in enumerate(zip(act_table.rows[0].cells, act_headers)):
            cell.width = act_widths[idx]
            _set_cell_background(cell, COLOR_NAVY_HEX)
            _set_cell_margins(cell, top=160, bottom=160, left=160, right=160)
            hp = cell.paragraphs[0]
            hr = hp.add_run(h_text)
            hr.font.name = "Arial"
            hr.font.size = Pt(9.5)
            hr.font.bold = True
            hr.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, act in enumerate(debrief.action_items):
            row = act_table.add_row()
            bg_color = COLOR_BG_ALT_HEX if (row_idx % 2 == 1) else COLOR_WHITE_HEX
            data = [f"☐  {act.task}", act.owner, act.deadline]

            for idx, (cell, val) in enumerate(zip(row.cells, data)):
                cell.width = act_widths[idx]
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                p = cell.paragraphs[0]
                run = p.add_run(val)
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.color.rgb = COLOR_SLATE_DARK
                if idx == 1:
                    run.font.bold = True
    else:
        p_none = doc.add_paragraph("No action items assigned.")
        p_none.runs[0].font.color.rgb = COLOR_SLATE_MUTED

    # 6. Spoken Transcript Appendix (if available)
    if transcript_history:
        doc.add_page_break()
        add_section_header("Appendix: Full Spoken Transcript Log")
        
        t_table = doc.add_table(rows=1, cols=3)
        t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(t_table)
        t_widths = [Inches(1.0), Inches(1.8), Inches(4.2)]

        for idx, (cell, h_text) in enumerate(zip(t_table.rows[0].cells, ["Time", "Identified Speaker", "Speech Content"])):
            cell.width = t_widths[idx]
            _set_cell_background(cell, COLOR_NAVY_HEX)
            _set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
            p = cell.paragraphs[0]
            hr = p.add_run(h_text)
            hr.font.name = "Arial"
            hr.font.size = Pt(9)
            hr.font.bold = True
            hr.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, chunk in enumerate(transcript_history):
            row = t_table.add_row()
            bg_color = COLOR_BG_ALT_HEX if (row_idx % 2 == 1) else COLOR_WHITE_HEX
            time_val = chunk.get("timestamp", "-")
            spk_val = chunk.get("speaker", "Guest")
            conf = chunk.get("confidence")
            if conf is not None:
                spk_val += f" ({conf}%)"
            text_val = chunk.get("text", "")

            for idx, (cell, val) in enumerate(zip(row.cells, [time_val, spk_val, text_val])):
                cell.width = t_widths[idx]
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                p = cell.paragraphs[0]
                run = p.add_run(val)
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.font.color.rgb = COLOR_SLATE_DARK

    # Save to in-memory bytes buffer
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
